"""工作台的两条验收口径（P0-04 WS-1 搜索、WS-5 预览）。

这两项此前主体已交付，但验收没达到：任务列表只按运行态分组，产物面板只有
文件名加下载按钮。
"""

from __future__ import annotations

import io

from openpyxl import Workbook

from cogniwork.core.config import get_settings

from .conftest import auth_header


def _prefix() -> str:
    return get_settings().api_prefix


def _xlsx() -> bytes:
    wb = Workbook()
    sheet = wb.active
    sheet.append(["Channel", "This week", "Last week"])
    sheet.append(["Email", 50, 40])
    sheet.append(["Paid social", 12, 30])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_task(client, headers, message: str) -> str:
    created = client.post(f"{_prefix()}/tasks", headers=headers, json={"message": message})
    assert created.status_code == 200, created.text
    return created.json()["id"]


def test_history_search_matches_title_and_original_request(client, registered):
    headers = auth_header(registered["token"])
    wanted = _make_task(client, headers, "Summarise the paid social results for Q3")
    _make_task(client, headers, "Draft the launch checklist")

    hits = client.get(f"{_prefix()}/tasks", headers=headers, params={"q": "paid social"}).json()
    assert [item["id"] for item in hits["tasks"]] == [wanted]

    # 大小写不敏感；用户记得的是自己怎么说的
    assert client.get(f"{_prefix()}/tasks", headers=headers, params={"q": "PAID"}).json()["tasks"]
    assert (
        client.get(f"{_prefix()}/tasks", headers=headers, params={"q": "nothing here"}).json()[
            "tasks"
        ]
        == []
    )
    # 空 q 等于不过滤，不该把列表清空
    assert len(client.get(f"{_prefix()}/tasks", headers=headers, params={"q": ""}).json()["tasks"])


def test_search_does_not_cross_accounts(client, registered):
    headers = auth_header(registered["token"])
    _make_task(client, headers, "My private quarterly numbers")
    other = client.post(
        f"{_prefix()}/auth/register",
        json={"email": "other@example.com", "password": "a-strong-password"},
    ).json()
    found = client.get(
        f"{_prefix()}/tasks",
        headers=auth_header(other["access_token"]),
        params={"q": "quarterly"},
    ).json()
    assert found["tasks"] == []


def _artifacts_for(client, headers, task_id) -> list[dict]:
    import time

    for _ in range(200):
        task = client.get(f"{_prefix()}/tasks/{task_id}", headers=headers).json()
        if task["status"] in {"succeeded", "failed", "cancelled", "timed_out"}:
            return task.get("artifacts") or []
        time.sleep(0.02)
    return []


def test_artifact_preview_renders_a_table_and_markdown(client, registered):
    headers = auth_header(registered["token"])
    uploaded = client.post(
        f"{_prefix()}/files",
        headers=headers,
        files={
            "file": (
                "numbers.xlsx",
                _xlsx(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    ).json()
    created = client.post(
        f"{_prefix()}/tasks",
        headers=headers,
        json={"message": "Turn this into a weekly report", "file_ids": [uploaded["id"]]},
    )
    task_id = created.json()["id"]
    artifacts = _artifacts_for(client, headers, task_id)
    assert artifacts, "零授权路径应产出至少一个产物"

    kinds = {}
    for item in artifacts:
        body = client.get(f"{_prefix()}/artifacts/{item['id']}/preview", headers=headers).json()
        kinds[item["filename"]] = body["preview"]["kind"]
        assert body["filename"] == item["filename"]

    # 口径是 xlsx/csv/md/docx/png 都能预览。周报路径产出 md 与 xlsx。
    assert "markdown" in kinds.values() or "table" in kinds.values(), kinds
    for filename, kind in kinds.items():
        if filename.endswith(".md"):
            assert kind == "markdown"
        if filename.endswith(".xlsx"):
            assert kind == "table"


def test_preview_unit_covers_the_five_declared_formats():
    from cogniwork.runtime.preview import build_preview

    table = build_preview("numbers.xlsx", "application/octet-stream", _xlsx())
    assert table["kind"] == "table"
    assert table["header"] == ["Channel", "This week", "Last week"]
    assert table["rows"][0] == ["Email", "50", "40"]
    assert table["row_count"] == 2

    csv_preview = build_preview("rows.csv", "text/csv", b"a,b\n1,2\n3,4\n")
    assert csv_preview["kind"] == "table"
    assert csv_preview["header"] == ["a", "b"]
    assert len(csv_preview["rows"]) == 2

    md = build_preview("report.md", "text/markdown", b"# Weekly report\n\nAll good.\n")
    assert md["kind"] == "markdown"
    assert "Weekly report" in md["text"]

    png = build_preview("chart.png", "image/png", b"\x89PNG\r\n\x1a\nfake")
    assert png["kind"] == "image"
    assert png["data_uri"].startswith("data:image/png;base64,")

    from docx import Document

    document = Document()
    document.add_paragraph("Quarterly summary")
    buf = io.BytesIO()
    document.save(buf)
    docx_preview = build_preview("summary.docx", "application/octet-stream", buf.getvalue())
    assert docx_preview["kind"] == "text"
    assert "Quarterly summary" in docx_preview["text"]


def test_unsupported_format_is_a_state_not_an_error():
    from cogniwork.runtime.preview import build_preview

    out = build_preview("archive.zip", "application/zip", b"PK\x03\x04")
    assert out == {"kind": "none", "reason": "unsupported_format"}


def test_large_tables_are_truncated_server_side():
    from cogniwork.runtime.preview import MAX_ROWS, build_preview

    body = "col\n" + "\n".join(str(i) for i in range(MAX_ROWS * 3))
    out = build_preview("big.csv", "text/csv", body.encode())
    assert out["truncated"] is True
    assert len(out["rows"]) == MAX_ROWS - 1
    assert out["row_count"] == MAX_ROWS * 3
